# python init

# ---------------------------------------
# external dependencies
# ---------------------------------------

import pdfminer
from pdfminer.pdfparser import PDFParser
from pdfminer.pdfdocument import PDFDocument
from pdfminer.pdfpage import PDFPage
from pdfminer.pdfpage import PDFTextExtractionNotAllowed
from pdfminer.pdfinterp import PDFResourceManager
from pdfminer.pdfinterp import PDFPageInterpreter
from pdfminer.pdfdevice import PDFDevice
from pdfminer.layout import LAParams
from pdfminer.converter import PDFPageAggregator

from PyPDF2 import PdfFileWriter, PdfFileReader

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.graphics.shapes import Rect
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.colors import PCMYKColor, PCMYKColorSep, Color, black, blue, red, white, green

import os
import sys
import re
import glob
import StringIO
import shutil
import time
import datetime
import warnings
import math 
from random import randint

from openpyxl import load_workbook
from openpyxl import Workbook

import numpy as np
import pandas as pd




# ---------------------------------------
# functions & methods
# ---------------------------------------

# pdfPositionHandling 
#----------------
class pdfPositionHandling:
 
    def parse_obj(self, lt_objs, pos_list, page_num):
    	position_list_new=pos_list

        # loop over the object list
        for obj in lt_objs:

            if isinstance(obj, pdfminer.layout.LTTextLine):
             	position_list_new.append([obj.bbox[0], obj.bbox[1], page_num,obj.get_text().replace('\n', '')])

            # if it's a textbox, also recurse
            if isinstance(obj, pdfminer.layout.LTTextBoxHorizontal):
                self.parse_obj(obj._objs, position_list_new, page_num)

            # if it's a container, recurse
            elif isinstance(obj, pdfminer.layout.LTFigure):
                self.parse_obj(obj._objs,position_list_new, page_num)

    def parsepdf(self, filename, startpage, endpage):

        # Open a PDF file.
        fp = open(filename, 'rb')

        # Create Position List
        position_list=[]

        # Create a PDF parser object associated with the file object.
        parser = PDFParser(fp)

        # Create a PDF document object that stores the document structure.
        # Password for initialization as 2nd parameter
        document = PDFDocument(parser)

        # Check if the document allows text extraction. If not, abort.
        if not document.is_extractable:
            raise PDFTextExtractionNotAllowed

        # Create a PDF resource manager object that stores shared resources.
        rsrcmgr = PDFResourceManager()

        # Create a PDF device object.
        device = PDFDevice(rsrcmgr)

        # BEGIN LAYOUT ANALYSIS
        # Set parameters for analysis.
        laparams = LAParams()

        # Create a PDF page aggregator object.
        device = PDFPageAggregator(rsrcmgr, laparams=laparams)

        # Create a PDF interpreter object.
        interpreter = PDFPageInterpreter(rsrcmgr, device)


        i = 0
        # loop over all pages in the document
        for page in PDFPage.create_pages(document):
            if i >= startpage and i <= endpage:
                # read the page into a layout object
                interpreter.process_page(page)
                layout = device.get_result()

                # extract text from this object
                # print(position_list)
                self.parse_obj(layout._objs,position_list, i)
                i += 1

        position_list = pd.DataFrame(position_list)
        position_list.columns=["pos_x", "pos_y", "page", "text"]

        return(position_list)

# flatten_list
#----------------
def flatten_list(container):
    for i in container:
        if isinstance(i, (list,tuple)):
            for j in flatten_list(i):
                yield j
        else:
            yield i
